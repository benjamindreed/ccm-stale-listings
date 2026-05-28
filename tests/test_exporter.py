import os
import tempfile
from openpyxl import load_workbook
from redfin_scout.models import Listing, ScoredListing
from redfin_scout.exporter import write_excel


def make_scored(score=75, dom=45, agent_email="bob@test.com", original_price=320000) -> ScoredListing:
    listing = Listing(
        mls_id="MLS001",
        address="789 Pine Rd, Tacoma, WA 98401",
        property_type="Condo",
        days_on_market=dom,
        current_price=300000,
        original_price=original_price,
        price_history=[
            {"date": "2024-01-01", "price": 320000},
            {"date": "2024-02-01", "price": 300000},
        ],
        remarks="",
        agent_name="Bob Jones",
        brokerage="Tacoma Realty",
        agent_id="agent2",
        agent_email=agent_email,
        agent_phone="253-555-1111",
        county="Pierce",
    )
    return ScoredListing(listing=listing, score=score, reasons=["45+ days on market", "Price reduced"])


def test_creates_xlsx_file():
    with tempfile.TemporaryDirectory() as d:
        path = os.path.join(d, "out.xlsx")
        write_excel([make_scored()], path)
        assert os.path.exists(path)


def test_has_13_columns():
    with tempfile.TemporaryDirectory() as d:
        path = os.path.join(d, "out.xlsx")
        write_excel([make_scored()], path)
        ws = load_workbook(path).active
        assert ws.max_column == 13


def test_header_row():
    with tempfile.TemporaryDirectory() as d:
        path = os.path.join(d, "out.xlsx")
        write_excel([make_scored()], path)
        ws = load_workbook(path).active
        headers = [ws.cell(1, c).value for c in range(1, 14)]
        assert headers[0] == "Priority Score"
        assert headers[1] == "Opportunity Reasons"
        assert headers[4] == "Days on Market"
        assert headers[12] == "County"


def test_data_row_values():
    with tempfile.TemporaryDirectory() as d:
        path = os.path.join(d, "out.xlsx")
        write_excel([make_scored(score=75)], path)
        ws = load_workbook(path).active
        assert ws.cell(2, 1).value == 75
        assert ws.cell(2, 3).value == "789 Pine Rd, Tacoma, WA 98401"
        assert ws.cell(2, 13).value == "Pierce"


def test_missing_fields_show_marker():
    with tempfile.TemporaryDirectory() as d:
        path = os.path.join(d, "out.xlsx")
        sl = make_scored(agent_email=None, original_price=None)
        write_excel([sl], path)
        ws = load_workbook(path).active
        assert ws.cell(2, 7).value == "[MISSING]"   # Original Price col
        assert ws.cell(2, 11).value == "[MISSING]"  # Agent Email col


def test_green_fill_score_60_plus():
    with tempfile.TemporaryDirectory() as d:
        path = os.path.join(d, "out.xlsx")
        write_excel([make_scored(score=65)], path)
        ws = load_workbook(path).active
        assert ws.cell(2, 1).fill.fgColor.rgb == "FF92D050"


def test_yellow_fill_score_30_to_59():
    with tempfile.TemporaryDirectory() as d:
        path = os.path.join(d, "out.xlsx")
        write_excel([make_scored(score=45)], path)
        ws = load_workbook(path).active
        assert ws.cell(2, 1).fill.fgColor.rgb == "FFFFEB9C"


def test_red_fill_score_below_30():
    with tempfile.TemporaryDirectory() as d:
        path = os.path.join(d, "out.xlsx")
        write_excel([make_scored(score=10)], path)
        ws = load_workbook(path).active
        assert ws.cell(2, 1).fill.fgColor.rgb == "FFFFC7CE"


def test_rows_sorted_by_score_descending():
    with tempfile.TemporaryDirectory() as d:
        path = os.path.join(d, "out.xlsx")
        write_excel([make_scored(score=30), make_scored(score=90), make_scored(score=60)], path)
        ws = load_workbook(path).active
        scores = [ws.cell(r, 1).value for r in range(2, 5)]
        assert scores == [90, 60, 30]
