import os
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font
from openpyxl.utils import get_column_letter
from redfin_scout.models import ScoredListing

HEADERS = [
    "Priority Score", "Opportunity Reasons", "Property Address", "Property Type",
    "Days on Market", "Current Price", "Original Price", "Price Reduction History",
    "Listing Agent", "Brokerage", "Agent Email", "Agent Phone", "County",
]
GREEN  = PatternFill("solid", fgColor="FF92D050")
YELLOW = PatternFill("solid", fgColor="FFFFEB9C")
RED    = PatternFill("solid", fgColor="FFFFC7CE")
MISSING = "[MISSING]"


def _price_history_str(history: list[dict]) -> str:
    if not history or len(history) <= 1:
        return MISSING
    return "; ".join(f"{h.get('date', '?')}→${int(h.get('price', 0)):,}" for h in history)


def write_excel(listings: list[ScoredListing], output_path: str) -> None:
    sorted_listings = sorted(listings, key=lambda sl: sl.score, reverse=True)
    date_str = os.path.basename(output_path).replace("redfin_scout_", "").replace(".xlsx", "")
    wb = Workbook()
    ws = wb.active
    ws.title = f"Redfin Scout {date_str}"

    for col, header in enumerate(HEADERS, start=1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = Font(bold=True)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(HEADERS))}1"

    for row_idx, sl in enumerate(sorted_listings, start=2):
        l = sl.listing
        row = [
            sl.score,
            sl.reasons_str,
            l.address or MISSING,
            l.property_type or MISSING,
            l.days_on_market,
            f"${l.current_price:,}" if l.current_price else MISSING,
            f"${l.original_price:,}" if l.original_price else MISSING,
            _price_history_str(l.price_history),
            l.agent_name or MISSING,
            l.brokerage or MISSING,
            l.agent_email or MISSING,
            l.agent_phone or MISSING,
            l.county or MISSING,
        ]
        for col_idx, value in enumerate(row, start=1):
            ws.cell(row=row_idx, column=col_idx, value=value)

        score_cell = ws.cell(row=row_idx, column=1)
        score_cell.fill = GREEN if sl.score >= 60 else (YELLOW if sl.score >= 30 else RED)

    col_widths = [10, 45, 40, 18, 8, 14, 14, 55, 25, 30, 28, 18, 14]
    for i, width in enumerate(col_widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width

    wb.save(output_path)
